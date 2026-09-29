from __future__ import annotations

from io import BytesIO
from pathlib import Path
import struct
import zlib
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from doc_to_md_converter import ConversionError, convert_docx
from doc_to_md_converter.cli import main


def png(rgb: tuple[int, int, int] = (255, 0, 0)) -> bytes:
    def chunk(name: bytes, data: bytes) -> bytes:
        content = name + data
        return struct.pack(">I", len(data)) + content + struct.pack(">I", zlib.crc32(content))
    raw = b"\x00" + bytes(rgb)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">2I5B", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def picture(paragraph, data: bytes) -> None:
    paragraph.add_run().add_picture(BytesIO(data), width=Inches(0.2))


def hyperlink(paragraph, text: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    node.append(run)
    paragraph._p.append(node)


def test_ordered_content_and_images_are_byte_identical(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("Методичка", level=1)
    paragraph = doc.add_paragraph("До таблицы: ")
    paragraph.add_run("полужирный").bold = True
    paragraph.add_run(" и ")
    paragraph.add_run("курсив").italic = True
    hyperlink(paragraph, "документ", "https://example.org/some page")
    doc.add_paragraph("Пункт", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Слева"
    table.cell(0, 1).text = "Справа"
    data = png()
    picture(table.cell(1, 0).paragraphs[0], data)
    table.cell(1, 1).text = "Значение | с разделителем"
    doc.add_paragraph("После таблицы")
    picture(doc.add_paragraph("Вот изображение: "), data)
    source = tmp_path / "study.docx"
    doc.save(source)

    logs: list[str] = []
    result = convert_docx(source, progress=logs.append)
    markdown = result.markdown_path.read_text(encoding="utf-8")
    assert markdown.index("# Методичка") < markdown.index("До таблицы") < markdown.index("| Слева") < markdown.index("После таблицы")
    assert "**полужирный**" in markdown
    assert "*курсив*" in markdown
    assert "[документ](<https://example.org/some%20page>)" in markdown
    assert "- Пункт" in markdown
    assert "Значение \\| с разделителем" in markdown
    assert markdown.count("![") == 2
    assert markdown.count("assets/images/image1.png") == 2
    assert result.images_count == 1  # same embedded bytes/relation reused
    assert (result.output_dir / "assets/images/image1.png").read_bytes() == data
    assert "Готово" in result.log_path.read_text(encoding="utf-8")
    assert any("Сохранён ресурс" in line for line in logs)


def test_multiple_images_header_and_no_overwrite(tmp_path: Path) -> None:
    doc = Document()
    red, blue = png(), png((0, 0, 255))
    picture(doc.add_paragraph(), red)
    picture(doc.sections[0].header.paragraphs[0], blue)
    source = tmp_path / "lesson.docx"
    doc.save(source)
    first = convert_docx(source)
    second = convert_docx(source)
    assert first.output_dir.name == "lesson_md"
    assert second.output_dir.name == "lesson_md_2"
    assert first.images_count == 2
    markdown = first.markdown_path.read_text(encoding="utf-8")
    assert "Колонтитулы" in markdown
    assert "Верхний колонтитул" in markdown
    assert "assets/images/image1.png" in markdown
    assert "assets/images/image2.png" in markdown
    assert {path.read_bytes() for path in (first.output_dir / "assets/images").iterdir()} == {red, blue}
    with pytest.raises(ConversionError, match="уже существует"):
        convert_docx(source, output_dir=first.output_dir)


def test_orphaned_media_is_saved_and_logged(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Текст методички")
    source = tmp_path / "orphan.docx"
    doc.save(source)
    extra = png((12, 34, 56))
    with ZipFile(source, "a") as archive:
        archive.writestr("word/media/unreferenced.png", extra)
    result = convert_docx(source)
    assert result.images_count == 1
    assert (result.output_dir / "assets/images/image1.png").read_bytes() == extra
    assert any("не встречен" in warning for warning in result.warnings)
    text = result.markdown_path.read_text(encoding="utf-8")
    assert "Текст методички" in text
    assert "Дополнительные изображения" in text
    assert "assets/images/image1.png" in text


def test_unsupported_drawing_has_visible_placeholder_and_warning(tmp_path: Path) -> None:
    doc = Document()
    node = OxmlElement("w:drawing")
    doc.add_paragraph().add_run()._r.append(node)
    source = tmp_path / "shape.docx"
    doc.save(source)
    result = convert_docx(source)
    assert "Неподдерживаемый графический объект" in result.markdown_path.read_text(encoding="utf-8")
    assert any("без извлекаемой картинки" in w for w in result.warnings)


def test_reject_invalid_inputs_without_creating_export(tmp_path: Path) -> None:
    bad = tmp_path / "text.docx"
    bad.write_text("not a zip", encoding="utf-8")
    with pytest.raises(ConversionError, match="Некорректный"):
        convert_docx(bad)
    assert not (tmp_path / "text_md").exists()
    with pytest.raises(ConversionError, match="расширением"):
        convert_docx(tmp_path / "wrong.txt")
    with pytest.raises(ConversionError, match="не найден"):
        convert_docx(tmp_path / "missing.docx")


def test_cli_writes_output_and_reports_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "from-cli.docx"
    doc = Document()
    doc.add_paragraph("CLI готов")
    doc.save(source)
    assert main([str(source)]) == 0
    assert "Готово:" in capsys.readouterr().out
    assert (tmp_path / "from-cli_md/index.md").exists()
    assert main([str(tmp_path / "missing.docx")]) == 1
    assert "Ошибка:" in capsys.readouterr().err


def test_vml_picture_is_extracted_from_legacy_markup(tmp_path: Path) -> None:
    doc = Document()
    image_data = png((9, 8, 7))
    ordinary = doc.add_paragraph()
    picture(ordinary, image_data)
    # Reuse the image rId in a legacy w:pict element.
    blip = ordinary._p.xpath(".//a:blip")[0]
    rid = blip.get(qn("r:embed"))
    legacy = doc.add_paragraph()
    pict = OxmlElement("w:pict")
    from lxml import etree
    v_image = etree.Element("{urn:schemas-microsoft-com:vml}imagedata")
    v_image.set(qn("r:id"), rid)
    pict.append(v_image)
    legacy.add_run()._r.append(pict)
    source = tmp_path / "legacy.docx"
    doc.save(source)

    result = convert_docx(source)
    assert result.images_count == 1
    assert result.markdown_path.read_text(encoding="utf-8").count("assets/images/image1.png") == 2
    assert (result.output_dir / "assets/images/image1.png").read_bytes() == image_data


def test_numbered_list_and_escaping(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("First", style="List Number")
    doc.add_paragraph("Second", style="List Number")
    doc.add_paragraph("# not a heading")
    doc.add_paragraph("Some [brackets] *star* and |pipe|")
    source = tmp_path / "list.docx"
    doc.save(source)
    result = convert_docx(source)
    text = result.markdown_path.read_text(encoding="utf-8")
    assert "1. First" in text and "1. Second" in text
    assert "\\# not a heading" in text
    assert r"Some \[brackets\] \*star\* and \|pipe\|" in text
